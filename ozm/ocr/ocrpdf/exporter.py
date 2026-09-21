# -*- coding: utf-8 -*-
"""Выгрузка результата: JSON, XLSX, CSV, отладочное изображение.

Структура JSON (обоснование выбора)
-----------------------------------
Единица истины — ЯЧЕЙКА со своим bbox, источником и confidence. Помимо этого
даётся «строковое» представление, удобное для расчётов. Такая двойная подача
выбрана вместо плоского списка строк потому, что:

  * merged cells (группа профиля, марка стали) физически принадлежат одной
    ячейке; в плоской строке они пришлось бы дублировать, теряя признак
    «значение получено из объединённой ячейки»;
  * трассируемость: у каждого значения сохраняются координаты на странице
    и, при распознавании по глифам, координаты и confidence каждого символа;
  * колонки «по элементам конструкции» задаются данными, а не схемой: их число
    и названия берутся из шапки листа.
"""
from __future__ import annotations

import json
import os
import re
from dataclasses import asdict


from .pdfbackend import fitz
from .models import Table, to_jsonable
from .structure import Column, LogicalRow, column_key, standards_of

LOW_CONF_FILL = "FFF2CC"      # ячейки на ручную проверку
FAILED_FILL = "FFC7CE"        # строки с непройденной арифметикой
HEADER_FILL = "DDEBF7"

# XML 1.0 / Excel: TAB, LF, CR можно; остальные управляющие ломают xlsx.
_ILLEGAL_XLSX_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def xlsx_safe(value):
    """Убрать из строки символы, из-за которых openpyxl бросает IllegalCharacterError."""
    if not isinstance(value, str):
        return value
    return _ILLEGAL_XLSX_CHARS.sub("", value)


# --------------------------------------------------------------------------- #
#  JSON
# --------------------------------------------------------------------------- #
def table_to_dict(table: Table, cols: list[Column], rows: list[LogicalRow],
                  all_glyphs: bool = False) -> dict:
    # внутри ключ колонки — индекс, наружу отдаём человекочитаемое название
    element_names = {column_key(c): (c.element or ("col%d" % c.index))
                     for c in cols if c.role == "element_mass"}

    def cell_ref(cell) -> dict:
        return {
            "text": cell.text,
            "normalized_value": cell.normalized_value,
            "value_kind": cell.value_kind,
            "page": cell.page,
            "bbox": [round(v, 2) for v in cell.bbox],
            "confidence": cell.confidence,
            "source": cell.source,
            "row": cell.row, "col": cell.col,
            "row_span": cell.row_span, "col_span": cell.col_span,
            "requires_review": cell.requires_review,
            "alt_text": getattr(cell, "alt_text", ""),
            "notes": cell.notes,
            "candidates": list(getattr(cell, "candidates", None) or []),
            # Посимвольная трассировка по умолчанию пишется только для ячеек,
            # отправленных на проверку: именно там она нужна человеку, а для
            # всей страницы это лишние ~1 МБ на таблицу (--all-glyphs включает).
            "glyphs": [
                {"char": g.char, "bbox": [round(v, 2) for v in g.bbox],
                 "confidence": g.confidence,
                 "alternatives": [[a, b] for a, b in g.alternatives]}
                for g in cell.glyphs
            ] if (all_glyphs or cell.requires_review) else [],
        }

    checks_by_row: dict[int, list] = {}
    for chk in table.checks:
        checks_by_row.setdefault(chk.row, []).append(to_jsonable(chk))

    out_rows = []
    for row in rows:
        if row.kind == "empty" and not row.cells:
            continue
        grade_cell = row.cells.get("steel_grade")
        grade, grade_std = standards_of(grade_cell)
        group_cell = row.cells.get("profile_group")
        group, group_std = standards_of(group_cell, bare=True)
        elements = {}
        for key, cell in row.cells.items():
            if key.startswith("element:"):
                elements[element_names.get(key, key.split(":", 1)[1])] = cell_ref(cell)
        checks = checks_by_row.get(row.row, [])
        status = "failed" if any(c["status"] == "failed" for c in checks) else (
            "ok" if any(c["status"] == "ok" for c in checks) else "not_checked")
        out_rows.append({
            "row": row.row,
            "row_kind": row.kind,
            "position": row.position,
            "profile_group": group,
            "profile_group_standards": group_std,
            "steel_grade": grade,
            "steel_grade_standards": grade_std,
            "profile_size": (row.cells["profile_size"].normalized_value
                             if "profile_size" in row.cells else None),
            # равновероятные прочтения марки («2011» → 20Б1 / 20Ш1 / 20К1) — для выбора в интерфейсе
            "profile_size_candidates": (list(getattr(row.cells["profile_size"], "candidates", None) or [])
                                        if "profile_size" in row.cells else []),
            "elements": elements,
            "total_mass_t": (row.cells["total_mass"].normalized_value
                             if "total_mass" in row.cells else None),
            "source_cells": {k: {"row": c.row, "col": c.col,
                                 "bbox": [round(v, 2) for v in c.bbox],
                                 "confidence": c.confidence, "source": c.source}
                             for k, c in row.cells.items()},
            "validation": status,
            "checks": checks,
            "notes": row.notes,
        })

    return {
        "index": table.index,
        "title": table.title,
        "kind": table.kind,
        "part": table.part,
        "notes": table.notes,
        "continues_table": table.continues_table,
        "page": table.page,
        "bbox": [round(v, 2) for v in table.bbox],
        "n_rows": table.n_rows,
        "n_cols": table.n_cols,
        "header_rows": table.header_rows,
        # единица масс, как подписана в шапке; значения в rows/cells — всегда тонны
        "mass_unit": next((c.unit for c in cols if c.role == "total_mass" and c.unit), "т"),
        "columns": [{"index": c.index, "letter": c.letter, "role": c.role,
                     "title": c.title, "element": c.element,
                     "unit": getattr(c, "unit", ""),
                     "bbox": [round(v, 2) for v in c.bbox]} for c in cols],
        "rows": out_rows,
        "cells": [cell_ref(c) for c in table.cells],
        "checks": [to_jsonable(c) for c in table.checks],
    }


def write_json(path: str, payload: dict) -> None:
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2)


# --------------------------------------------------------------------------- #
#  XLSX
# --------------------------------------------------------------------------- #
def write_xlsx(path: str, table: Table, cols: list[Column],
               rows: list[LogicalRow]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Таблица"
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="center", horizontal="center")

    # 1) лист «как на чертеже»: сетка с объединёнными ячейками
    from openpyxl.cell.cell import MergedCell
    for cell in table.cells:
        r, c = cell.row + 1, cell.col + 1
        xc = ws.cell(row=r, column=c)
        # Страховка: если сетка всё же оказалась противоречивой, пропускаем
        # ячейку, а не падаем на записи в объединённую область.
        if isinstance(xc, MergedCell):
            continue
        if cell.value_kind in ("number", "int") and cell.normalized_value is not None:
            xc.value = cell.normalized_value
            xc.number_format = "0.00" if cell.value_kind == "number" else "0"
        else:
            text = xlsx_safe(cell.text)
            xc.value = text or None
        xc.alignment = wrap
        xc.border = border
        if cell.row in table.header_rows:
            xc.fill = PatternFill("solid", fgColor=HEADER_FILL)
            xc.font = Font(bold=True)
        elif cell.requires_review:
            xc.fill = PatternFill("solid", fgColor=LOW_CONF_FILL)
        if cell.row_span > 1 or cell.col_span > 1:
            ws.merge_cells(start_row=r, start_column=c,
                           end_row=r + cell.row_span - 1,
                           end_column=c + cell.col_span - 1)
    for c in range(1, table.n_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 16

    # 2) плоский лист для расчётов
    ws2 = wb.create_sheet("Плоская")
    element_cols = [c for c in cols if c.role == "element_mass"]
    elements = [c.element or ("col%d" % c.index) for c in element_cols]
    head = (["№ строки", "Тип строки", "№ п.п.", "Наименование профиля",
             "Марка металла", "ГОСТ/ТУ", "Номер или размеры"]
            + elements + ["Общая масса, т", "Проверка", "Мин. confidence"])
    ws2.append([xlsx_safe(h) for h in head])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = wrap
    checks_failed = {c.row for c in table.checks if c.status == "failed"}
    for row in rows:
        if row.kind == "empty" and not row.cells:
            continue
        grade, std = standards_of(row.cells.get("steel_grade"))
        group, gstd = standards_of(row.cells.get("profile_group"), bare=True)
        vals = [row.row, row.kind, row.position, group, grade,
                "; ".join(std or gstd),
                (row.cells["profile_size"].normalized_value
                 if "profile_size" in row.cells else None)]
        # Число, которое не удалось разобрать («0.14 5», «12,4?»), раньше
        # уходило в плоский лист ПУСТОЙ ячейкой — распознанный текст терялся
        # именно там, где пользователь его ищет. Теперь остаётся текст.
        def _flat(cell):
            if cell is None:
                return None
            if cell.normalized_value is not None:
                return cell.normalized_value
            return cell.text.strip() or None
        for col in element_cols:
            vals.append(_flat(row.cells.get(column_key(col))))
        vals.append(_flat(row.cells.get("total_mass")))
        vals.append("failed" if row.row in checks_failed else "ok")
        confs = [c.confidence for c in row.cells.values() if not c.is_empty]
        vals.append(round(min(confs), 3) if confs else None)
        ws2.append([xlsx_safe(v) for v in vals])
        if row.row in checks_failed:
            for cell in ws2[ws2.max_row]:
                cell.fill = PatternFill("solid", fgColor=FAILED_FILL)
    for c in range(1, len(head) + 1):
        ws2.column_dimensions[get_column_letter(c)].width = 18

    # 3) лист проверок арифметики
    ws3 = wb.create_sheet("Проверки")
    ws3.append(["№ строки", "Проверка", "Итог из PDF", "Пересчитано",
                "Разница", "Допуск", "Статус", "Причина", "Слагаемые"])
    for cell in ws3[1]:
        cell.font = Font(bold=True)
    for chk in table.checks:
        ws3.append([xlsx_safe(v) for v in [
            chk.row, chk.kind, chk.expected, chk.calculated, chk.delta,
            chk.tolerance, chk.status, chk.reason,
            ", ".join(str(t) for t in chk.terms)]])
    for c in range(1, 10):
        ws3.column_dimensions[get_column_letter(c)].width = 16
    wb.save(path)


def write_csv(path: str, table: Table) -> None:
    """Сетка «как на чертеже» в CSV (для быстрого визуального сравнения)."""
    import csv
    grid = [["" for _ in range(table.n_cols)] for _ in range(table.n_rows)]
    for cell in table.cells:
        grid[cell.row][cell.col] = cell.text.replace("\n", " ")
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        csv.writer(fh, delimiter=";").writerows(grid)


# --------------------------------------------------------------------------- #
#  отладочное изображение
# --------------------------------------------------------------------------- #
def write_debug_png(path: str, page: fitz.Page, blocks, tables: list[Table],
                    dpi: int = 130) -> None:
    """Страница с прорисованными рамками таблиц, ячейками и проблемными местами.

    Цвета: синий — рамка блока, серый — ячейка, зелёный — merged-ячейка,
    оранжевый — ячейка на ручную проверку, красный — строка, не прошедшая
    арифметическую проверку.
    """
    import numpy as np
    import cv2

    pix = page.get_pixmap(dpi=dpi)
    img = np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.width, pix.n)
    img = cv2.cvtColor(img[:, :, :3], cv2.COLOR_RGB2BGR).copy()
    z = dpi / 72.0

    def rect(bbox, color, th=1):
        cv2.rectangle(img, (int(bbox[0] * z), int(bbox[1] * z)),
                      (int(bbox[2] * z), int(bbox[3] * z)), color, th)

    failed_rows = {(t.index, c.row) for t in tables for c in t.checks
                   if c.status == "failed"}
    for blk in blocks:
        rect(blk.bbox, (255, 128, 0), 3)
    for t in tables:
        for cell in t.cells:
            if (t.index, cell.row) in failed_rows:
                rect(cell.bbox, (0, 0, 255), 2)
            elif cell.requires_review:
                rect(cell.bbox, (0, 165, 255), 2)
            elif cell.row_span > 1 or cell.col_span > 1:
                rect(cell.bbox, (0, 170, 0), 1)
            else:
                rect(cell.bbox, (170, 170, 170), 1)
        cv2.putText(img, "table_%d %s" % (t.index, t.kind),
                    (int(t.bbox[0] * z), max(12, int(t.bbox[1] * z) - 6)),
                    cv2.FONT_HERSHEY_SIMPLEX, 0.5, (255, 0, 0), 1, cv2.LINE_AA)
    legend = [("table block", (255, 128, 0)), ("cell", (170, 170, 170)),
              ("merged cell", (0, 170, 0)), ("needs review", (0, 165, 255)),
              ("arithmetic failed", (0, 0, 255))]
    y = 18
    for name, color in legend:
        cv2.rectangle(img, (8, y - 9), (24, y + 2), color, -1)
        cv2.putText(img, name, (30, y), cv2.FONT_HERSHEY_SIMPLEX, 0.45,
                    (0, 0, 0), 1, cv2.LINE_AA)
        y += 18
    cv2.imwrite(path, img)


def write_report(path: str, report: dict) -> None:
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(report, fh, ensure_ascii=False, indent=2)
